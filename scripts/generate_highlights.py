import docx
from docx.shared import Inches, Pt

def generate_highlights():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run('Research Highlights')
    r_title.bold = True
    r_title.font.size = Pt(13)
    p_title.paragraph_format.space_after = Pt(10)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run('Title: "Ontology-Constrained Neuro-Symbolic Decoding: Enforcing Axiomatic Least-Privilege in Autonomous LLM Agents"')
    r_sub.italic = True
    p_sub.paragraph_format.space_after = Pt(12)

    bullets = [
        "Formal RBAC ontologies compile into runtime token grammars for LLM agents.",
        "Logit masking mathematically guarantees zero privilege escalations.",
        "O-CTD suppresses attack success rate from 80% to 20% in enterprise tests.",
        "Achieves 100% benign task completion with zero JSON syntax formatting errors.",
        "Wilcoxon signed-rank test confirms statistical significance (p < 0.001)."
    ]

    for b in bullets:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(b)
        r.font.size = Pt(11)

    doc.save('paper/highlights.docx')
    doc.save('E:/Downloads/highlights.docx')

    with open('paper/highlights.txt', 'w', encoding='utf-8') as f:
        f.write("Research Highlights (Knowledge-Based Systems)\n")
        f.write("Note: Each bullet is strictly under Elsevier's 85-character limit (including spaces).\n\n")
        for b in bullets:
            f.write(f"• {b} ({len(b)} characters)\n")

    print("[SUCCESS] Successfully generated paper/highlights.docx and E:/Downloads/highlights.docx")

if __name__ == "__main__":
    generate_highlights()
